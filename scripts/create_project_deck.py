"""Create an editable PowerPoint and matching shareable PDF overview."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "attached_assets"
W, H, SCALE = 1600, 900, 120
NAVY = "#101c2a"
PANEL = "#1d2d3c"
PANEL2 = "#273a4b"
WHITE = "#f8f4ec"
MUTED = "#b6c2cd"
RED = "#d84755"
GOLD = "#f1d28b"
GREEN = "#9dd9bf"
LINE = "#455667"
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def rgb(color):
    return RGBColor.from_string(color.lstrip("#"))


class Slide:
    def __init__(self, deck, number, section):
        self.image = Image.new("RGB", (W, H), NAVY)
        self.draw = ImageDraw.Draw(self.image)
        self.slide = deck.slides.add_slide(deck.slide_layouts[6])
        self.rect(0, 0, W, H, NAVY)
        self.txt(76, 48, "SAINT PHILIP NERI", 21, GOLD, True)
        self.txt(1535, 48, section.upper(), 17, MUTED, True, "right")
        self.line(76, 822, 1524, 822, LINE, 2)
        self.txt(76, 846, "SCHOOL PAPERWORK  /  PROJECT OVERVIEW",
                 15, MUTED)
        self.txt(1524, 846, f"{number:02d}  /  08", 15, MUTED, False, "right")

    def rect(self, x, y, w, h, color, rounded=0, border=None):
        xy = (round(x), round(y), round(x + w), round(y + h))
        self.draw.rounded_rectangle(xy, radius=rounded, fill=color,
                                    outline=border, width=2 if border else 1)
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(kind, Inches(x / SCALE),
                     Inches(y / SCALE), Inches(w / SCALE), Inches(h / SCALE))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background() if border is None else None
        if border:
            shape.line.color.rgb = rgb(border)
            shape.line.width = Pt(1.2)

    def line(self, x1, y1, x2, y2, color, width=2):
        self.draw.line((x1, y1, x2, y2), fill=color, width=width)
        line = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                 Inches(x1 / SCALE), Inches(y1 / SCALE),
                 Inches((x2 - x1) / SCALE),
                 Inches(max(width, y2 - y1) / SCALE))
        line.fill.solid()
        line.fill.fore_color.rgb = rgb(color)
        line.line.fill.background()

    def txt(self, x, y, value, size=30, color=WHITE, bold=False,
            align="left", width=None):
        f = ImageFont.truetype(BOLD if bold else FONT, size)
        lines = value.split("\n")
        for i, part in enumerate(lines):
            left = x
            if align != "left":
                length = self.draw.textlength(part, font=f)
                left = x - (length if align == "right" else length / 2)
            self.draw.text((left, y + i * size * 1.45), part, fill=color, font=f)
        # Use a separate native text box for each line so PowerPoint keeps it editable.
        for i, part in enumerate(lines):
            box_x = x
            line_width = width or (W - x - 50)
            if align != "left":
                estimated = max(30, self.draw.textlength(part, font=f) + 24)
                line_width = estimated
                box_x = x - (estimated if align == "right" else estimated / 2)
            tb = self.slide.shapes.add_textbox(Inches(box_x / SCALE),
                   Inches((y + i * size * 1.45 - 2) / SCALE),
                   Inches(line_width / SCALE),
                   Inches((size * 1.5) / SCALE))
            tf = tb.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = 0
            tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = part
            p.font.name = "DejaVu Sans"
            p.font.bold = bold
            p.font.size = Pt(size * 0.6)
            p.font.color.rgb = rgb(color)

    def circle(self, x, y, r, color):
        self.draw.ellipse((x - r, y - r, x + r, y + r), fill=color)
        sh = self.slide.shapes.add_shape(MSO_SHAPE.OVAL,
                     Inches((x-r) / SCALE), Inches((y-r) / SCALE),
                     Inches(2*r / SCALE), Inches(2*r / SCALE))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(color)
        sh.line.fill.background()

    def logo(self, x, y, size):
        path = ROOT / "static" / "logo.jpg"
        if not path.exists():
            return
        img = Image.open(path).convert("RGB")
        img.thumbnail((size, size))
        px, py = int(x + (size - img.width) / 2), int(y + (size - img.height) / 2)
        self.image.paste(img, (px, py))
        self.slide.shapes.add_picture(str(path), Inches(px / SCALE),
              Inches(py / SCALE), width=Inches(img.width / SCALE),
              height=Inches(img.height / SCALE))

    def heading(self, eyebrow, title, subtitle):
        self.txt(76, 127, eyebrow.upper(), 21, RED, True)
        self.txt(76, 174, title, 56, WHITE, True)
        self.txt(79, 262, subtitle, 25, MUTED)

    def card(self, x, y, width, height, marker, title, lines, accent=GOLD):
        self.rect(x, y, width, height, PANEL, 18, LINE)
        self.rect(x + 25, y + 33, 7, 42, accent, 3)
        self.txt(x + 52, y + 36, marker.upper(), 19, accent, True)
        self.txt(x + 26, y + 105, title, 29, WHITE, True)
        for i, line in enumerate(lines):
            self.txt(x + 26, y + 170 + i * 42, line, 21, MUTED)


def make_deck():
    OUT.mkdir(exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(W/SCALE), Inches(H/SCALE)
    pages = []

    s = Slide(deck, 1, "Project introduction")
    s.rect(70, 148, 9, 505, RED)
    s.txt(112, 166, "School paperwork,", 64, WHITE, True)
    s.txt(112, 258, "made simpler.", 64, GOLD, True)
    s.txt(115, 402, "A connected space for parents to complete forms,", 29, MUTED)
    s.txt(115, 453, "find volunteer work and track service hours.", 29, MUTED)
    s.rect(1105, 182, 318, 318, PANEL, 30)
    s.logo(1162, 239, 205)
    s.txt(115, 610, "PARENT EXPERIENCE  +  ADMIN TOOLS", 22, GREEN, True)
    pages.append(s.image)

    s = Slide(deck, 2, "The purpose")
    s.heading("The need", "A simpler way to stay connected",
              "One school process, shared by families and staff.")
    s.card(76, 351, 455, 349, "01  /  Forms", "Know what is due",
           ["Review required paperwork", "and submit it online."], GOLD)
    s.card(572, 351, 455, 349, "02  /  Volunteering", "Find ways to help",
           ["See available jobs and", "sign up when eligible."], GREEN)
    s.card(1068, 351, 455, 349, "03  /  Time", "Keep a clear record",
           ["Log service sessions for", "admin verification."], RED)
    pages.append(s.image)

    s = Slide(deck, 3, "Parent experience")
    s.heading("For parents", "Everything starts at Dashboard",
              "Four focused sections keep the parent experience organized.")
    s.rect(76, 345, 1448, 350, PANEL, 20, LINE)
    for i, (name, desc, color) in enumerate([
        ("My Forms", "Pending and submitted paperwork", GOLD),
        ("Jobs", "Volunteer opportunities", GREEN),
        ("Labor Hours", "Sessions and service history", RED),
        ("Profile", "Name and classroom details", GOLD)
    ]):
        x = 105 + i*358
        s.rect(x, 385, 327, 80, PANEL2, 15)
        s.txt(x + 20, 405, name, 28, color, True)
        s.txt(x + 8, 520, desc.split(" and ")[0], 18, MUTED)
        if " and " in desc:
            s.txt(x + 8, 552, "and " + desc.split(" and ")[1], 18, MUTED)
    s.txt(112, 635, "Top navigation: Dashboard  /  Notifications  /  Parent name", 20, MUTED)
    pages.append(s.image)

    s = Slide(deck, 4, "Forms")
    s.heading("Parent workflow", "Paperwork without the chase",
              "Parents can see what needs attention and what is already done.")
    s.card(76, 348, 692, 354, "Pending forms", "Review & submit",
           ["See active forms and deadlines.", "Open a form and submit it online."], GOLD)
    s.card(802, 348, 722, 354, "Submitted forms", "Revisit past work",
           ["Check submission details.", "View or edit a submitted form."], GREEN)
    pages.append(s.image)

    s = Slide(deck, 5, "Volunteer jobs")
    s.heading("Parent workflow", "Turn interest into involvement",
              "Jobs, signups and eligibility all live in the same system.")
    labels = [
        ("01", "Explore", "Browse active school jobs."),
        ("02", "Qualify", "Eligibility is checked for roles\nworking with children."),
        ("03", "Join", "Sign up for opportunities\nthat fit.")
    ]
    for i, (n, title, desc) in enumerate(labels):
        x = 76 + i*494
        s.rect(x, 361, 455, 319, PANEL, 18, LINE)
        s.txt(x+26, 388, n, 30, RED, True)
        s.txt(x+26, 458, title, 37, WHITE, True)
        s.txt(x+26, 548, desc, 20, MUTED)
    pages.append(s.image)

    s = Slide(deck, 6, "Labor hours")
    s.heading("Parent + admin workflow", "From work session to verified hours",
              "A simple sequence makes volunteer time easier to document.")
    for i, (num, title, sub, color) in enumerate([
        ("01", "Check in", "Start a session", GOLD),
        ("02", "Check out", "Record time worked", GREEN),
        ("03", "Verify", "Admin confirms work", RED)
    ]):
        x = 76 + i*494
        s.card(x, 360, 455, 304, num, title, [sub], color)
    s.txt(79, 723, "Parents can review history and print or export service logs.",
          22, MUTED)
    pages.append(s.image)

    s = Slide(deck, 7, "Admin tools")
    s.heading("For administrators", "The tools behind the experience",
              "Manage requests, eligibility and service records from one place.")
    entries = [
        ("Overview", "Forms, submissions and school activity"),
        ("Applications", "Review volunteer job applications"),
        ("Verify Parents", "Manage parent job eligibility"),
        ("Labor Verification", "Review sessions and confirm hours")
    ]
    for i, (title, detail) in enumerate(entries):
        y = 344 + i*97
        s.rect(76, y, 1448, 77, PANEL, 13, LINE)
        s.circle(114, y+39, 13, RED if i==3 else GOLD)
        s.txt(155, y+15, title, 26, WHITE, True)
        s.txt(1490, y+22, detail, 19, MUTED, False, "right")
    s.txt(79, 751, "Also available: job postings and all-parent service-hour summaries.",
          20, GREEN)
    pages.append(s.image)

    s = Slide(deck, 8, "A connected experience")
    s.circle(800, 223, 72, RED)
    s.txt(771, 170, "✓", 86, WHITE, True)
    s.txt(800, 355, "Less paperwork.", 67, WHITE, True, "center")
    s.txt(800, 458, "More time for school.", 67, GOLD, True, "center")
    s.txt(800, 601, "Forms  •  Jobs  •  Labor hours  •  Verification",
          27, MUTED, False, "center")
    s.txt(800, 689, "SAINT PHILIP NERI  /  SCHOOL PAPERWORK",
          19, GREEN, True, "center")
    pages.append(s.image)

    pptx = OUT / "saint_philip_neri_project_overview.pptx"
    pdf = OUT / "saint_philip_neri_project_overview.pdf"
    deck.save(pptx)
    pages[0].save(pdf, "PDF", resolution=120, save_all=True,
                  append_images=pages[1:])
    print(f"{pptx}\n{pdf}")


if __name__ == "__main__":
    make_deck()